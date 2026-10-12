"""
Reporte completo: resumen mensual de inadmitidos / internaciones en Excel.

Replica la hoja "<Aerolínea>_Mensual" que se armaba a mano a partir de las
hojas diarias: (1) tabla por mes y nacionalidad, (2) Inadmitidos vs
Internaciones por país y (3) totales de pasajeros por mes, más dos gráficas
nativas de Excel. Usa las mismas definiciones que el reporte diario
(_compute_inadmitidos_data) para que las cifras cuadren con él.
"""
from datetime import datetime, timezone as dt_timezone
from io import BytesIO
from types import SimpleNamespace
from urllib.parse import quote

from django.db.models import Count, Q
from django.http import HttpResponse

from apps.cuentas.roles import flujo_principal_required

from ..models import Registro
from ..utils.excel_estilos import (
    MAROON, DARK_RED, ROSA_TITULO, ROSA_FECHAS, VERDE_INTERN, DATA_BG, PAX_BG,
)
from .registros import MESES_ES, _detectar_vuelo_y_origen, _filtros_arribos

# Colores de las series en las gráficas (el verde de la tabla es demasiado
# claro para barras, aquí va oscurecido).
COLOR_SERIE_INAD = '6B1D1D'
COLOR_SERIE_INTERN = '4F7A65'
COLOR_SERIE_SR = '4D060A'


def _rango_mes_utc(anio, mes):
    """Ventana [1º del mes, 1º del mes siguiente) en UTC, igual que las
    ventanas por día del reporte diario."""
    inicio = datetime(anio, mes, 1, tzinfo=dt_timezone.utc)
    if mes == 12:
        fin = datetime(anio + 1, 1, 1, tzinfo=dt_timezone.utc)
    else:
        fin = datetime(anio, mes + 1, 1, tzinfo=dt_timezone.utc)
    return inicio, fin


def _ordenar_por_total(por_pais):
    """[(país, valores_por_mes)] de mayor a menor total; empate alfabético."""
    return sorted(por_pais.items(), key=lambda kv: (-sum(kv[1]), kv[0]))


def _compute_reporte_mensual(anio, meses):
    """Agrega por mes los conteos del reporte de inadmitidos.

    Dos consultas por mes (totales + desglose por país) en lugar de las ~10
    por día de la ruta diaria."""
    filtro_tij, filtro_mex = _filtros_arribos()
    filtro_arribos = filtro_tij | filtro_mex
    n = len(meses)

    totales = {k: [] for k in ('inad', 'intern', 'sr', 'local', 'transito', 'total')}
    inad_por_pais, intern_por_pais = {}, {}

    for i, mes in enumerate(meses):
        inicio, fin = _rango_mes_utc(anio, mes)
        arribos = Registro.objects.filter(
            vuelo_fecha__gte=inicio,
            vuelo_fecha__lt=fin,
        ).filter(filtro_arribos)

        tot = arribos.aggregate(
            inad=Count('id', filter=Q(segunda_revision=True, rechazado=True), distinct=True),
            intern=Count('id', filter=Q(segunda_revision=True, internacion=True), distinct=True),
            sr=Count('id', filter=Q(segunda_revision=True), distinct=True),
            local=Count('id', filter=filtro_tij, distinct=True),
            transito=Count('id', filter=filtro_mex, distinct=True),
            total=Count('id', distinct=True),
        )
        for clave, lista in totales.items():
            lista.append(tot[clave] or 0)

        por_pais = (
            arribos.filter(segunda_revision=True)
            .order_by()
            .values('pais_emision')
            .annotate(
                inad=Count('id', filter=Q(rechazado=True), distinct=True),
                intern=Count('id', filter=Q(internacion=True), distinct=True),
            )
        )
        for item in por_pais:
            nat = item['pais_emision'] or 'Sin especificar'
            if item['inad']:
                inad_por_pais.setdefault(nat, [0] * n)[i] += item['inad']
            if item['intern']:
                intern_por_pais.setdefault(nat, [0] * n)[i] += item['intern']

    # Vuelo y origen: primer arribo de toda la selección.
    inicio_sel, _ = _rango_mes_utc(anio, meses[0])
    _, fin_sel = _rango_mes_utc(anio, meses[-1])
    arribos_sel = Registro.objects.filter(
        vuelo_fecha__gte=inicio_sel,
        vuelo_fecha__lt=fin_sel,
    ).filter(filtro_arribos).distinct()
    vuelo, _ciudad, origen = _detectar_vuelo_y_origen(arribos_sel)

    return {
        'anio': anio,
        'meses': list(meses),
        'meses_labels': [MESES_ES[m].capitalize() for m in meses],
        'vuelo': vuelo,
        'origen': origen,
        **totales,
        'inad_por_pais': _ordenar_por_total(inad_por_pais),
        'intern_por_pais': _ordenar_por_total(intern_por_pais),
    }


def _estilos():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin = Side(style='thin', color='FFE2E2E5')
    return SimpleNamespace(
        fill_maroon=PatternFill('solid', fgColor=MAROON),
        fill_dark_red=PatternFill('solid', fgColor=DARK_RED),
        fill_rosa_tit=PatternFill('solid', fgColor=ROSA_TITULO),
        fill_rosa_fec=PatternFill('solid', fgColor=ROSA_FECHAS),
        fill_verde=PatternFill('solid', fgColor=VERDE_INTERN),
        fill_data=PatternFill('solid', fgColor=DATA_BG),
        fill_pax=PatternFill('solid', fgColor=PAX_BG),
        # Tipografía Calibri en todo (igual que el Excel diario)
        font_data=Font(name='Calibri', size=11, color='FF1A1C1E'),
        font_blanco=Font(name='Calibri', size=11, bold=True, color='FFFFFFFF'),
        font_maroon=Font(name='Calibri', size=11, bold=True, color='FF6B1D1D'),
        font_fechas=Font(name='Calibri', size=10, bold=True, color='FF554241'),
        font_pax=Font(name='Calibri', size=11, bold=True, color='FF1A1C1E'),
        font_intern=Font(name='Calibri', size=11, bold=True, color='FF1C2522'),
        font_tot_inad=Font(name='Calibri', size=12, bold=True, color='FF6B1D1D'),
        font_tot_intern=Font(name='Calibri', size=12, bold=True, color='FF1C2522'),
        font_tot_sr=Font(name='Calibri', size=12, bold=True, color='FFFFFFFF'),
        thin=thin,
        box=Side(style='thin', color='FF000000'),
        borde=Border(left=thin, right=thin, top=thin, bottom=thin),
        center=Alignment(horizontal='center', vertical='center', wrap_text=True),
        left=Alignment(horizontal='left', vertical='center', wrap_text=True),
        indent=Alignment(horizontal='left', vertical='center', indent=2),
    )


def _escribir_fila(ws, est, r, c0, ancho, valores, fill, font,
                   label_fill=None, label_font=None, label_align=None):
    """Escribe `ancho` celdas desde (r, c0). La primera es la etiqueta y puede
    llevar su propio relleno / fuente / alineación."""
    for j in range(ancho):
        valor = valores[j] if j < len(valores) else None
        celda = ws.cell(row=r, column=c0 + j, value=valor)
        es_label = j == 0
        celda.fill = label_fill if (es_label and label_fill is not None) else fill
        celda.font = label_font if (es_label and label_font is not None) else font
        celda.border = est.borde
        celda.alignment = (label_align or est.left) if es_label else est.center
        if isinstance(valor, int):
            celda.number_format = '#,##0'


def _enmarcar(ws, est, r, c0, ancho):
    """Borde negro alrededor de la fila (verticales internos quedan tenues)."""
    from openpyxl.styles import Border

    for j in range(ancho):
        izq = est.box if j == 0 else est.thin
        der = est.box if j == ancho - 1 else est.thin
        ws.cell(row=r, column=c0 + j).border = Border(left=izq, right=der, top=est.box, bottom=est.box)


def _ancho_cm(ws, col_ini, col_fin):
    """Ancho aproximado en cm de un rango de columnas (para dimensionar la
    gráfica sin que invada la tabla de la derecha)."""
    from openpyxl.utils import get_column_letter

    total = 0
    for c in range(col_ini, col_fin + 1):
        w = ws.column_dimensions[get_column_letter(c)].width or 8.43
        total += w * 0.19
    return total


def _escribir_reporte_mensual(ws, data):
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.utils import get_column_letter

    est = _estilos()
    labels = data['meses_labels']
    n = len(labels)

    # La columna Total sólo aporta con varios meses.
    incluir_total = n > 1
    last_col = (n + 2) if incluir_total else (n + 1)  # etiqueta + n meses [+ Total]

    def con_total(etiqueta, valores):
        fila = [etiqueta] + list(valores)
        return fila + [sum(valores)] if incluir_total else fila

    # ─── Tabla 1 · por mes y nacionalidad ───
    r = 1
    if last_col >= 4:
        _escribir_fila(ws, est, r, 1, last_col, ['Vuelo:', data['vuelo'], 'Origen:', data['origen']],
                       est.fill_maroon, est.font_blanco)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=last_col)
        ws.cell(row=r, column=4).alignment = est.left
    else:
        _escribir_fila(ws, est, r, 1, last_col, [f"Vuelo: {data['vuelo']}   Origen: {data['origen']}"],
                       est.fill_maroon, est.font_blanco)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=last_col)

    r += 1
    # Como texto: con '#,##0' el año saldría "2.026".
    _escribir_fila(ws, est, r, 1, last_col, [str(data['anio'])], est.fill_rosa_tit, est.font_maroon)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=last_col)
    ws.cell(row=r, column=1).alignment = est.center

    r += 1
    _escribir_fila(ws, est, r, 1, last_col, ['Nacionalidad'], est.fill_maroon, est.font_blanco)
    r += 1
    fila_encabezado = r  # aquí arranca también la tabla 2
    encabezado = [None] + labels + (['Total'] if incluir_total else [])
    _escribir_fila(ws, est, r, 1, last_col, encabezado, est.fill_rosa_fec, est.font_fechas)

    for nat, vals in data['inad_por_pais']:
        r += 1
        _escribir_fila(ws, est, r, 1, last_col, con_total(nat, vals), est.fill_data, est.font_data,
                       label_fill=est.fill_rosa_tit)
        _enmarcar(ws, est, r, 1, last_col)

    r += 1
    _escribir_fila(ws, est, r, 1, last_col, con_total('Total Inadmitidos', data['inad']),
                   est.fill_rosa_tit, est.font_tot_inad)
    _enmarcar(ws, est, r, 1, last_col)
    r += 1
    _escribir_fila(ws, est, r, 1, last_col, con_total('Total Internaciones:', data['intern']),
                   est.fill_verde, est.font_tot_intern)
    _enmarcar(ws, est, r, 1, last_col)

    for nat, vals in data['intern_por_pais']:
        r += 1
        _escribir_fila(ws, est, r, 1, last_col, con_total(nat, vals), est.fill_data, est.font_data,
                       label_fill=est.fill_verde, label_align=est.indent)
        _enmarcar(ws, est, r, 1, last_col)

    r += 1
    _escribir_fila(ws, est, r, 1, last_col, con_total('Total Segundas Revisiones:', data['sr']),
                   est.fill_dark_red, est.font_tot_sr)
    r += 1
    _escribir_fila(ws, est, r, 1, last_col, con_total('Local:', data['local']), est.fill_data, est.font_data)
    r += 1
    _escribir_fila(ws, est, r, 1, last_col, con_total('En tránsito:', data['transito']), est.fill_data, est.font_data)
    r += 1
    _escribir_fila(ws, est, r, 1, last_col, con_total('Total pasajeros:', data['total']), est.fill_pax, est.font_pax)
    fin_tabla1 = r

    # ─── Tabla 3 · resumen por mes (debajo de la tabla 1) ───
    r = fin_tabla1 + 3
    fila_t3 = r
    _escribir_fila(ws, est, r, 1, last_col, encabezado, est.fill_rosa_fec, est.font_fechas)
    r += 1
    _escribir_fila(ws, est, r, 1, last_col, con_total('Inadmitidos', data['inad']),
                   est.fill_rosa_tit, est.font_tot_inad)
    r += 1
    _escribir_fila(ws, est, r, 1, last_col, con_total('Internaciones', data['intern']),
                   est.fill_verde, est.font_tot_intern)
    r += 1
    _escribir_fila(ws, est, r, 1, last_col, con_total('Total Segundas Rev.', data['sr']),
                   est.fill_dark_red, est.font_tot_sr)
    r += 1
    _escribir_fila(ws, est, r, 1, last_col, con_total('Total pasajeros:', data['total']), est.fill_pax, est.font_pax)
    fin_tabla3 = r

    # ─── Tabla 2 · Inadmitidos vs Internaciones por país (a la derecha) ───
    c2 = last_col + 2  # una columna de separación
    inad_tot = {nat: sum(vals) for nat, vals in data['inad_por_pais']}
    intern_tot = {nat: sum(vals) for nat, vals in data['intern_por_pais']}
    con_inad = sorted(inad_tot, key=lambda p: (-inad_tot[p], -intern_tot.get(p, 0), p))
    solo_intern = sorted((p for p in intern_tot if p not in inad_tot), key=lambda p: (-intern_tot[p], p))

    r = fila_encabezado
    encabezados_t2 = (
        ('País', est.fill_maroon, est.font_blanco),
        ('Inadmitidos', est.fill_rosa_tit, est.font_maroon),
        ('Internaciones', est.fill_verde, est.font_intern),
        ('Total', est.fill_dark_red, est.font_blanco),
    )
    for j, (texto, fill, font) in enumerate(encabezados_t2):
        celda = ws.cell(row=r, column=c2 + j, value=texto)
        celda.fill, celda.font, celda.border, celda.alignment = fill, font, est.borde, est.center

    for pais in con_inad + solo_intern:
        r += 1
        i_val, n_val = inad_tot.get(pais, 0), intern_tot.get(pais, 0)
        es_solo_intern = pais not in inad_tot
        # Los ceros quedan en blanco, como en la hoja original.
        _escribir_fila(ws, est, r, c2, 4, [pais, i_val or None, n_val or None, i_val + n_val],
                       est.fill_data, est.font_pax, label_font=est.font_data,
                       label_fill=est.fill_verde if es_solo_intern else est.fill_rosa_tit,
                       label_align=est.indent if es_solo_intern else None)
    fin_paises = r
    r += 1
    total_inad, total_intern = sum(inad_tot.values()), sum(intern_tot.values())
    _escribir_fila(ws, est, r, c2, 4, ['Total', total_inad, total_intern, total_inad + total_intern],
                   est.fill_pax, est.font_pax)
    _enmarcar(ws, est, r, c2, 4)
    fin_tabla2 = r

    # ─── Anchos de columna ───
    ws.column_dimensions['A'].width = 26
    for c in range(2, last_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 10
    ws.column_dimensions[get_column_letter(c2 - 1)].width = 3
    ws.column_dimensions[get_column_letter(c2)].width = 18
    for c in range(c2 + 1, c2 + 4):
        ws.column_dimensions[get_column_letter(c)].width = 13

    # ─── Gráficas nativas ───
    def _etiquetas_valor():
        dl = DataLabelList()
        dl.showVal = True
        dl.showSerName = dl.showCatName = dl.showLegendKey = dl.showPercent = False
        return dl

    def _pintar(chart, colores):
        for serie, color in zip(chart.series, colores):
            serie.graphicalProperties.solidFill = color
            serie.graphicalProperties.line.solidFill = color

    # 1) Segundas revisiones por mes (sin la columna Total ni pasajeros, por escala).
    graf_mes = BarChart()
    graf_mes.type = 'col'
    graf_mes.grouping = 'clustered'
    graf_mes.title = 'Pasajeros y Segundas Revisiones'
    graf_mes.add_data(Reference(ws, min_col=1, max_col=n + 1, min_row=fila_t3 + 1, max_row=fila_t3 + 3),
                      from_rows=True, titles_from_data=True)
    graf_mes.set_categories(Reference(ws, min_col=2, max_col=n + 1, min_row=fila_t3, max_row=fila_t3))
    _pintar(graf_mes, (COLOR_SERIE_INAD, COLOR_SERIE_INTERN, COLOR_SERIE_SR))
    graf_mes.dataLabels = _etiquetas_valor()
    graf_mes.legend.position = 'b'
    graf_mes.x_axis.delete = False
    graf_mes.y_axis.delete = False
    graf_mes.width = max(9, _ancho_cm(ws, 1, c2 - 1))
    graf_mes.height = 8
    ws.add_chart(graf_mes, f'A{fin_tabla3 + 2}')

    # 2) Segunda revisión por país (barras horizontales, sin la fila Total).
    if fin_paises > fila_encabezado:
        graf_pais = BarChart()
        graf_pais.type = 'bar'
        graf_pais.grouping = 'clustered'
        graf_pais.title = 'Pasajeros en Segunda Revisión por País'
        graf_pais.add_data(Reference(ws, min_col=c2 + 1, max_col=c2 + 2,
                                     min_row=fila_encabezado, max_row=fin_paises),
                           titles_from_data=True)
        graf_pais.set_categories(Reference(ws, min_col=c2, min_row=fila_encabezado + 1, max_row=fin_paises))
        _pintar(graf_pais, (COLOR_SERIE_INAD, COLOR_SERIE_INTERN))
        graf_pais.dataLabels = _etiquetas_valor()
        graf_pais.legend.position = 'b'
        graf_pais.x_axis.scaling.orientation = 'maxMin'  # el país con más casos arriba
        graf_pais.x_axis.delete = False
        graf_pais.y_axis.delete = False
        graf_pais.width = 16
        graf_pais.height = max(8, 0.55 * (fin_paises - fila_encabezado) + 3)
        ws.add_chart(graf_pais, f'{get_column_letter(c2)}{fin_tabla2 + 2}')

    # Impresión: horizontal, ajustado a una página de ancho.
    ws.page_setup.orientation = 'landscape'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def _sufijo_meses(meses):
    """'sep' · 'may-sep' (consecutivos) · 'may_jul_sep' (salteados)."""
    abrevs = [MESES_ES[m] for m in meses]
    if len(meses) == 1:
        return abrevs[0]
    if meses == list(range(meses[0], meses[-1] + 1)):
        return f'{abrevs[0]}-{abrevs[-1]}'
    return '_'.join(abrevs)


@flujo_principal_required
def generar_reporte_mensual_excel(request):
    """Descarga el Reporte completo (.xlsx): ?anio=2026&meses=5,6,7,8,9"""
    from openpyxl import Workbook

    try:
        anio = int(request.GET.get('anio', ''))
        meses = sorted({int(m) for m in request.GET.get('meses', '').split(',') if m.strip()})
    except ValueError:
        return HttpResponse('Parámetros inválidos: use anio=AAAA y meses=1,2,…', status=400)

    if not 2000 <= anio <= 2100:
        return HttpResponse('Año inválido', status=400)
    if not meses:
        return HttpResponse('Selecciona al menos un mes', status=400)
    if any(m < 1 or m > 12 for m in meses):
        return HttpResponse('Los meses deben estar entre 1 y 12', status=400)

    data = _compute_reporte_mensual(anio, meses)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Resumen mensual'
    _escribir_reporte_mensual(ws, data)

    buffer = BytesIO()
    wb.save(buffer)
    xlsx_bytes = buffer.getvalue()
    buffer.close()

    nombre = f"Reporte_Completo_{data['vuelo']}_{anio}_{_sufijo_meses(meses)}.xlsx"
    response = HttpResponse(
        xlsx_bytes,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f"attachment; filename*=UTF-8''{quote(nombre)}"
    return response
