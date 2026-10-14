from datetime import date, datetime, timezone as dt_timezone
from io import BytesIO

from django.contrib.auth.models import Group, User
from django.test import TestCase
from django.urls import reverse

from apps.cuentas.roles import GRUPO_AEROPUERTO, GRUPO_GENERAL
from .models import Registro, UploadBatch
from .utils.parsers import obtener_nacionalidad
from .views.registros import _compute_inadmitidos_data
from .views.reporte_mensual import _compute_reporte_mensual


def _registro(batch, fecha, llegada, pais, sr=False, r=False, pi=False):
    return Registro.objects.create(
        batch=batch, vuelo_numero='HU7925', vuelo_fecha=fecha,
        aeropuerto_salida='PEK', aeropuerto_llegada=llegada,
        nombre_pasajero='PASAJERO', numero_asiento='1A', estado_checkin='OK',
        numero_ticket='T', genero='M', codigo_pais_emision='XX', pais_emision=pais,
        segunda_revision=sr, rechazado=r, internacion=pi,
    )


class ReporteMensualTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.aeropuerto = User.objects.create_user('aero', password='x')
        cls.aeropuerto.groups.add(Group.objects.get(name=GRUPO_AEROPUERTO))
        cls.general = User.objects.create_user('gen', password='x')
        cls.general.groups.add(Group.objects.get(name=GRUPO_GENERAL))

        tij = UploadBatch.objects.create(usuario=cls.aeropuerto, vuelo_numero='HU7925', tipo_vuelo='PEK-TIJ')
        mex = UploadBatch.objects.create(usuario=cls.aeropuerto, vuelo_numero='HU7925', tipo_vuelo='PEK-MEX')
        regreso = UploadBatch.objects.create(usuario=cls.aeropuerto, vuelo_numero='HU7926')

        mayo = datetime(2026, 5, 20, tzinfo=dt_timezone.utc)
        # Último instante de mayo en UTC: debe contar en mayo, no en junio.
        fin_mayo = datetime(2026, 5, 31, 23, 59, tzinfo=dt_timezone.utc)
        junio = datetime(2026, 6, 3, tzinfo=dt_timezone.utc)

        # Mayo: 2 inadmitidos (China, Bangladés), 1 internación (México), 1 SR sin resolver, 1 sin SR
        _registro(tij, mayo, 'TIJ', 'China', sr=True, r=True)
        _registro(mex, mayo, 'MEX', 'Bangladés', sr=True, r=True)
        _registro(mex, mayo, 'MEX', 'México', sr=True, pi=True)
        _registro(tij, fin_mayo, 'TIJ', 'China', sr=True)
        _registro(mex, mayo, 'MEX', 'China')
        # Junio: Bangladés vuelve a aparecer (debe ser la misma fila), 1 internación China
        _registro(tij, junio, 'TIJ', 'Bangladés', sr=True, r=True)
        _registro(mex, junio, 'MEX', 'China', sr=True, pi=True)
        # Vuelo de regreso MEX→PEK: no es arribo, no debe contar
        Registro.objects.create(
            batch=regreso, vuelo_numero='HU7926', vuelo_fecha=junio,
            aeropuerto_salida='MEX', aeropuerto_llegada='PEK',
            nombre_pasajero='REGRESO', numero_asiento='1A', estado_checkin='OK',
            numero_ticket='T', genero='M', codigo_pais_emision='CN', pais_emision='China',
            segunda_revision=True, rechazado=True,
        )

    def test_conteos_por_mes(self):
        d = _compute_reporte_mensual(2026, [5, 6])
        self.assertEqual(d['meses_labels'], ['May', 'Jun'])
        self.assertEqual(d['inad'], [2, 1])
        self.assertEqual(d['intern'], [1, 1])
        # La SR sin resolver de fin de mayo no cuenta: SR = Inadmitidos + Internaciones
        self.assertEqual(d['sr'], [3, 2])
        self.assertEqual(d['local'], [2, 1])
        self.assertEqual(d['transito'], [3, 1])
        self.assertEqual(d['total'], [5, 2])
        self.assertEqual(d['vuelo'], 'HU7925')
        self.assertEqual(d['origen'], 'China')

    def test_nacionalidad_repetida_en_varios_meses_es_una_sola_fila(self):
        d = _compute_reporte_mensual(2026, [5, 6])
        self.assertEqual(d['inad_por_pais'], [('Bangladés', [1, 1]), ('China', [1, 0])])
        self.assertEqual(d['intern_por_pais'], [('China', [0, 1]), ('México', [1, 0])])

    def test_parametros_invalidos(self):
        self.client.force_login(self.aeropuerto)
        url = reverse('generar_reporte_mensual_excel')
        for params in ({}, {'anio': '2026'}, {'anio': 'x', 'meses': '5'}, {'anio': '2026', 'meses': '0,13'}):
            self.assertEqual(self.client.get(url, params).status_code, 400, params)

    def test_usuario_general_no_accede(self):
        self.client.force_login(self.general)
        resp = self.client.get(reverse('generar_reporte_mensual_excel'), {'anio': '2026', 'meses': '5'})
        self.assertEqual(resp.status_code, 302)

    def test_descarga_xlsx(self):
        from openpyxl import load_workbook

        self.client.force_login(self.aeropuerto)
        resp = self.client.get(reverse('generar_reporte_mensual_excel'), {'anio': '2026', 'meses': '5,6'})
        self.assertEqual(resp.status_code, 200)
        self.assertIn('Reporte_Completo_HU7925_2026_may-jun.xlsx', resp['Content-Disposition'])

        ws = load_workbook(BytesIO(resp.content))['Resumen mensual']
        filas = {ws.cell(row=r, column=1).value: [ws.cell(row=r, column=c).value for c in range(2, 5)]
                 for r in range(1, ws.max_row + 1)}
        self.assertEqual(filas['Total Inadmitidos'], [2, 1, 3])
        self.assertEqual(filas['Total pasajeros:'], [5, 2, 7])
        self.assertEqual(len(ws._charts), 2)

    def test_pagina_expone_meses_disponibles(self):
        self.client.force_login(self.aeropuerto)
        resp = self.client.get(reverse('inadmitidos_page'))
        self.assertEqual(resp.context['meses_disponibles'], {2026: [5, 6]})
        self.assertContains(resp, 'id="btn-reporte-completo"')

    def test_meses_sin_vuelos_no_se_ofrecen(self):
        # Ruta suspendida jul-sep: agosto tiene datos, julio y septiembre no.
        tij = UploadBatch.objects.get(tipo_vuelo='PEK-TIJ')
        _registro(tij, datetime(2026, 8, 3, 7, tzinfo=dt_timezone.utc), 'TIJ', 'China')
        _registro(tij, datetime(2026, 12, 31, 23, tzinfo=dt_timezone.utc), 'TIJ', 'China')
        _registro(tij, datetime(2027, 1, 2, 7, tzinfo=dt_timezone.utc), 'TIJ', 'China')
        self.client.force_login(self.aeropuerto)
        resp = self.client.get(reverse('inadmitidos_page'))
        self.assertEqual(resp.context['meses_disponibles'], {2026: [5, 6, 8, 12], 2027: [1]})

    def test_sr_diaria_cuenta_sin_estado(self):
        # El reporte del día sí cuenta las SR en curso (aún sin R ni I).
        d = _compute_inadmitidos_data(date(2026, 5, 20), date(2026, 5, 31))
        self.assertEqual(d['totals_sr'][0], 3)   # 20-may: 2 R + 1 I
        self.assertEqual(d['totals_sr'][-1], 1)  # 31-may: sólo la SR sin resolver

    def test_pagina_avisa_sr_sin_estado(self):
        self.client.force_login(self.aeropuerto)
        resp = self.client.get(reverse('inadmitidos_page'))
        pendientes = resp.context['sr_sin_estado']
        self.assertEqual(len(pendientes), 1)
        self.assertEqual(pendientes[0]['vuelo_fecha'], datetime(2026, 5, 31, 23, 59, tzinfo=dt_timezone.utc))
        # La fecha se muestra en UTC (31/05), igual que las ventanas del reporte.
        self.assertContains(resp, 'id="sr-sin-estado"')
        self.assertContains(resp, '31/05/2026')
        self.assertNotContains(resp, 'toggleSrSinEstado(this)')  # con 3 o menos no hay "Leer más"

        # Con más de 3 sólo los 3 primeros quedan a la vista; el resto tras "Leer más".
        tij = UploadBatch.objects.get(tipo_vuelo='PEK-TIJ')
        for _ in range(4):
            _registro(tij, datetime(2026, 6, 10, tzinfo=dt_timezone.utc), 'TIJ', 'China', sr=True)
        resp = self.client.get(reverse('inadmitidos_page'))
        self.assertEqual(len(resp.context['sr_sin_estado']), 5)
        self.assertContains(resp, 'class="sr-extra hidden"', count=2)
        self.assertContains(resp, 'Leer más (2)', count=2)  # texto del botón + data-mas

        Registro.objects.filter(segunda_revision=True, rechazado=False, internacion=False).update(internacion=True)
        resp = self.client.get(reverse('inadmitidos_page'))
        self.assertEqual(resp.context['sr_sin_estado'], [])
        self.assertNotContains(resp, 'id="sr-sin-estado"')


class NacionalidadesTests(TestCase):
    def test_territorios_y_codigos_icao(self):
        self.assertEqual(obtener_nacionalidad('MAC'), 'Macao')
        self.assertEqual(obtener_nacionalidad('hkg'), 'Hong Kong')
        self.assertEqual(obtener_nacionalidad('TWN'), 'Taiwán')
        self.assertEqual(obtener_nacionalidad('GBN'), 'Reino Unido (BNO)')
        self.assertEqual(obtener_nacionalidad('XX'), 'XX')


class AdminListNacionalidadTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.aeropuerto = User.objects.create_user('aero', password='x')
        cls.aeropuerto.groups.add(Group.objects.get(name=GRUPO_AEROPUERTO))
        cls.lote = UploadBatch.objects.create(usuario=cls.aeropuerto, vuelo_numero='HU7925', tipo_vuelo='PEK-TIJ')
        otro = UploadBatch.objects.create(usuario=cls.aeropuerto, vuelo_numero='HU7925', tipo_vuelo='PEK-MEX')
        fecha = datetime(2026, 5, 20, tzinfo=dt_timezone.utc)
        _registro(cls.lote, fecha, 'TIJ', 'Estados Unidos')
        _registro(otro, fecha, 'MEX', 'Estados Unidos')
        _registro(cls.lote, fecha, 'TIJ', 'México')
        _registro(cls.lote, fecha, 'TIJ', '')

    def setUp(self):
        self.client.force_login(self.aeropuerto)

    def test_lista_nacionalidades_con_total(self):
        resp = self.client.get(reverse('admin_list'))
        self.assertEqual(resp.context['nacionalidades'], [
            {'pais_emision': 'Estados Unidos', 'total': 2},
            {'pais_emision': 'México', 'total': 1},
        ])
        self.assertContains(resp, 'id="nacionalidadesData"')

    def test_filtra_por_pais_de_emision(self):
        resp = self.client.get(reverse('admin_list'), {'pais': 'Estados Unidos'})
        self.assertEqual({r.pais_emision for r in resp.context['page_obj']}, {'Estados Unidos'})
        self.assertEqual(len(resp.context['page_obj']), 2)

    def test_lote_por_enlace_sigue_filtrando(self):
        resp = self.client.get(reverse('admin_list'), {'batch': self.lote.id, 'pais': 'Estados Unidos'})
        self.assertEqual(len(resp.context['page_obj']), 1)
        self.assertEqual(resp.context['lote_activo'], self.lote)
        self.assertContains(resp, 'Quitar filtro de lote')
        # Un batch no numérico se ignora en vez de romper la página
        self.assertEqual(self.client.get(reverse('admin_list'), {'batch': 'abc'}).status_code, 200)
